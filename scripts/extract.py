import openpyxl, json, datetime, os
from collections import defaultdict, Counter

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
REPO_ROOT = os.path.dirname(SCRIPT_DIR)
NEW_SRC = os.path.join(REPO_ROOT, 'data', 'raw.xlsx')
OLD_SRC = os.path.join(REPO_ROOT, 'reference', 'master.xlsx')

TEAMS = ['FS WEST', 'FS EAST', 'FS Central', 'FS South']

# ---------- reference tables (still sourced from the original master workbook) ----------
wb_old = openpyxl.load_workbook(OLD_SRC, read_only=True, data_only=True)

ws_m = wb_old['Members']
member_team = {}
team_members = defaultdict(list)
for row in ws_m.iter_rows(min_row=1, max_row=96, values_only=True):
    a = row[0]
    if a in TEAMS:
        name = row[1]
        if name and name not in team_members[a]:
            team_members[a].append(name)
            member_team[name] = a

ws_ig = wb_old['ITEM Group']
model_to_group = {}
for row in ws_ig.iter_rows(min_row=1, max_row=98, values_only=True):
    model, group = row[1], row[3]
    if model and group:
        model_to_group[model] = group

# 김민결: excluded from the dashboard per user request
EXCLUDED_PEOPLE = {'김민결'}
for _p in EXCLUDED_PEOPLE:
    if _p in member_team:
        _t = member_team.pop(_p)
        if _p in team_members.get(_t, []):
            team_members[_t].remove(_p)

# account -> team map, sourced from the master workbook's '거래처 List' sheet (H column,
# 담당 팀). These are plain stored values (not formulas), so unlike the per-team revisit
# tabs they are NOT affected by the FS EAST dropdown/reference bug.
ws_al = wb_old['거래처 List']
account_team_map = {}
for row in ws_al.iter_rows(min_row=3, max_row=5048, values_only=True):
    acct, team = row[0], row[7]
    if acct and team in TEAMS:
        account_team_map[acct] = team

GROUPS_ORDER = ['CBC','COA','HISCL','Urine','A1c','RF-500']

BUCKET = {}
for p in ['고장 수리']: BUCKET[p] = '고장수리'
for p in ['PM', '온라인 정기점검']: BUCKET[p] = 'PM'
for p in ['장비 설치', '장비 이전 설치', '설치/데모 평가', '설치', '장비 철수', '철수 및 폐기']: BUCKET[p] = '설치이전'
for p in ['Calibration', '원격 Calibration', 'Evaluation', 'EVALUATION', '데이터 문제 해결 (QC/검체)', 'Quality Control']: BUCKET[p] = 'Cal평가'
def bucket(p):
    return BUCKET.get(p, '기타')

# ---------- KPI weighting table (provided by user, "기술 요청" applies to 서비스요청,
# "학술 요청" applies to 학술요청 form-type rows) ----------
TECH_WEIGHT = {
    '고장 수리':1, 'PM':0.4, 'GMCS':0.2, '부품':0.05, '장비 이전 설치':1, '장비 설치':2, '소모품':1,
    '장비 철수':1, '기타':1, 'Calibration':1, '데이터 TroubleShooting':1, 'EVALUATION':2, 'Evaluation':2,
    '전산':1, '교육':1, '설치/데모 평가':2, '인증 지원':1, '데이터 문제 해결 (QC/검체)':1,
    '사용자 교육, PT발표':1, '사용자 교육, PT 발표':1, '임상 연구 지원':1, '고객 학술 지원':1, '고객 관리':1,
    '장비 troubleShooting':1, '장비 TroubleShooting':1, 'Accreditation Support':1, 'Data Troubleshooting':1,
    '원격A/S':1, 'Quality Control':1, 'WAM':1, '온라인 정기점검':0.2, '원격 Calibration':10, '설치':0.1,
    '고객 요구 사항 지원':0.1, '방문 및 관계 형성':0.1, 'Interface':0.1, 'H/P 개발':0.1,
}
ACAD_WEIGHT = {
    '설치/데모 평가':2, 'Calibration':1, '인증 지원':1, '데이터 문제 해결 (QC/검체)':1,
    '사용자 교육, PT 발표':1, '사용자 교육, PT발표':1, '임상 연구 지원':1, '고객 학술 지원':1, '고객 관리':1,
    'EVALUATION':2, 'Evaluation':2, '장비 troubleShooting':1, '장비 TroubleShooting':1,
    'Accreditation Support':1, 'Data Troubleshooting':1, '데이터 TroubleShooting':1, 'Quality Control':1,
    'WAM':1, 'Training':1, 'Proficiency Test':1, 'Etc':0.05, 'Report':0.05, '기타':1,
}
def kpi_weight(formtype, purpose):
    if purpose is None:
        return 1.0  # small number of rows with no resolvable purpose text; treat as generic (기타=1) work
    if formtype == '학술 요청':
        return ACAD_WEIGHT.get(purpose, 1.0)
    # default: treat 서비스 요청 and any other form type via the technical table
    return TECH_WEIGHT.get(purpose, 1.0)

# ---------- fiscal year (Sysmex Korea: Apr 1 - Mar 31) ----------
TODAY = datetime.datetime.now()
def fy_start_year(d):
    return d.year if d.month >= 4 else d.year - 1

CUR_FY = fy_start_year(TODAY)
PREV_FY = CUR_FY - 1

def period_of(d):
    fy = fy_start_year(d)
    if fy == CUR_FY: return '이번 년도'
    if fy == PREV_FY: return '직전 년도'
    return '전체 누적'

PERIODS = ['이번 년도','직전 년도','전체 누적']
FY_LABEL = {
    '이번 년도': 'FY%d (%d.04–%d.03)' % (CUR_FY, CUR_FY, CUR_FY+1),
    '직전 년도': 'FY%d (%d.04–%d.03)' % (PREV_FY, PREV_FY, PREV_FY+1),
    '전체 누적': 'FY%d 이전' % PREV_FY,
}
print('TODAY', TODAY, 'CUR_FY', CUR_FY, 'PREV_FY', PREV_FY)
print('FY_LABEL', FY_LABEL)

# ---------- parse new raw data ----------
wb = openpyxl.load_workbook(NEW_SRC, read_only=True, data_only=True)
ws = wb['2. 모든 약속_활동']

def resolve_purpose(subj, purpose_col):
    if purpose_col is not None:
        return purpose_col
    if not subj:
        return None
    parts = subj.split('/')
    if len(parts) >= 4:
        return '/'.join(parts[3:])
    if len(parts) == 2:
        return parts[1]
    return None

def model_from_subject(subj):
    if not subj:
        return None
    parts = subj.split('/')
    if len(parts) >= 3:
        return parts[1]
    return None

all_rows = [row for row in ws.iter_rows(min_row=2, values_only=True) if row is not None and row[5] is not None]
print('total data rows loaded:', len(all_rows))

# pass 1: build a Serial No. -> most-common-Model lookup, to recover rows whose
# subject text doesn't carry the model (e.g. "거래처/목적" only, 2 segments)
serial_to_model_counter = defaultdict(Counter)
for row in all_rows:
    subj = row[8]; serial = row[14]
    m = model_from_subject(subj)
    if m and serial:
        serial_to_model_counter[serial][m] += 1
serial_to_model = {s: c.most_common(1)[0][0] for s, c in serial_to_model_counter.items()}
print('serial->model lookup built:', len(serial_to_model), 'serials')

def resolve_model(subj, serial):
    m = model_from_subject(subj)
    if m:
        return m
    if serial and serial in serial_to_model:
        return serial_to_model[serial]
    return None

person_group = defaultdict(lambda: defaultdict(Counter))
person_mttr = defaultdict(lambda: defaultdict(lambda: defaultdict(lambda: {'sum':0.0,'n':0})))
person_case = defaultdict(lambda: defaultdict(Counter))
person_month = defaultdict(lambda: defaultdict(Counter))
person_account = defaultdict(Counter)
person_score = defaultdict(lambda: defaultdict(float))               # [person][period] -> total weighted score
person_score_group = defaultdict(lambda: defaultdict(lambda: defaultdict(float)))  # [person][period][group] -> score

n = 0
excluded_cancelled = 0
excluded_mttr_outlier = 0
skipped_group = 0
unmatched_purpose_rows = 0

for row in all_rows:
    n += 1
    status = row[3]
    if status == '취소됨':
        excluded_cancelled += 1
        continue
    st_dt = row[5]; en_dt = row[6]
    person = row[7]
    subj = row[8]
    account = row[12]
    purpose_col = row[16]

    if not person or person in EXCLUDED_PEOPLE:
        continue
    if not isinstance(st_dt, datetime.datetime):
        continue

    period = period_of(st_dt)
    formtype = row[4]
    purpose = resolve_purpose(subj, purpose_col)
    model = resolve_model(subj, row[14])
    group = model_to_group.get(model) if model else None
    if group not in GROUPS_ORDER:
        group = None

    w = kpi_weight(formtype, purpose)
    person_score[person][period] += w
    if group:
        person_score_group[person][period][group] += w

    if group:
        person_group[person][period][group] += 1
        if purpose == '고장 수리' and isinstance(en_dt, datetime.datetime):
            mins = (en_dt - st_dt).total_seconds() / 60.0
            if 0 <= mins <= 500:
                person_mttr[person][period][group]['sum'] += mins
                person_mttr[person][period][group]['n'] += 1
            elif mins > 500:
                excluded_mttr_outlier += 1
    else:
        skipped_group += 1

    b = bucket(purpose)
    if purpose is None:
        unmatched_purpose_rows += 1
    person_case[person][period][b] += 1
    person_month[person][period][st_dt.month] += 1

    if account and b in ('고장수리', 'PM'):
        person_account[person][account] += 1

print('rows processed:', n)
print('excluded cancelled:', excluded_cancelled)
print('excluded mttr outlier (>500min):', excluded_mttr_outlier)
print('skipped group (#N/A / unmapped model):', skipped_group)
print('rows with no resolvable purpose:', unmatched_purpose_rows)

# ---------- revisit rate / FTFR ----------
# Reverse-engineered from the master workbook's '거래처 List' + 'FSG 모든 직원' formulas
# (읽어보면 AD/AE 헬퍼 열): for each piece of equipment (Serial No.), a "고장 수리" (repair)
# visit counts as a REVISIT if the same serial had another 고장수리 or PM visit in the
# 14 days before it; it counts against FTFR specifically if the same serial+Error1 had
# another 고장수리 visit in the 14 days before it. This flag is computed globally across
# all teams (equipment doesn't belong to a team, accounts do), then rolled up per account
# and the account is attributed to a team via the master 거래처 List mapping. Applying this
# one formula uniformly to every team (rather than each team's own precomputed sheet) is
# what fixes FS EAST, whose own sheet formula was broken (a bad dropdown reference produced
# #N/A instead of real numbers).
revisit_recs = []
for row in all_rows:
    if row[3] == '취소됨':
        continue
    st_dt = row[5]; en_dt = row[6]
    subj = row[8]
    account = row[12]
    serial = row[14]
    err1 = row[9]
    purpose = resolve_purpose(subj, row[16])
    if not isinstance(en_dt, datetime.datetime) or not serial or not account:
        continue
    err_norm = err1.strip().lower() if isinstance(err1, str) else err1
    revisit_recs.append({
        'account': account, 'serial': serial, 'err': err_norm, 'end': en_dt,
        'purpose': purpose, 'period': period_of(st_dt) if isinstance(st_dt, datetime.datetime) else None,
    })

by_serial_r = defaultdict(list)
for rec in revisit_recs:
    by_serial_r[rec['serial']].append(rec)
for lst in by_serial_r.values():
    lst.sort(key=lambda x: x['end'])
    for i, rec in enumerate(lst):
        if rec['purpose'] != '고장 수리':
            rec['AD'] = None; rec['AE'] = None
            continue
        window_start = rec['end'] - datetime.timedelta(days=14)
        rec['AD'] = 1 if any(o is not rec and window_start <= o['end'] < rec['end'] and o['purpose'] in ('고장 수리', 'PM') for o in lst) else 2
        rec['AE'] = 1 if any(o is not rec and window_start <= o['end'] < rec['end'] and o['purpose'] == '고장 수리' and o['err'] == rec['err'] for o in lst) else 2

account_recs = defaultdict(list)
for rec in revisit_recs:
    account_recs[rec['account']].append(rec)

REVISIT_PERIODS = ['이번 년도', '직전 년도', 'cum']
def account_rates(acct):
    evs = account_recs.get(acct, [])
    out = {}
    for per in REVISIT_PERIODS:
        sub = evs if per == 'cum' else [e for e in evs if e['period'] == per]
        denom_rv = sum(1 for e in sub if e['purpose'] in ('고장 수리', 'PM'))
        numer_rv = sum(1 for e in sub if e.get('AD') == 1)
        denom_ftfr = sum(1 for e in sub if e['purpose'] == '고장 수리')
        numer_bad = sum(1 for e in sub if e.get('AE') == 1)
        out[per] = {
            'revisit': (numer_rv / denom_rv) if denom_rv else None,
            'ftfr': ((denom_ftfr - numer_bad) / denom_ftfr) if denom_ftfr else None,
            'n_rv': denom_rv, 'n_ftfr': denom_ftfr,
        }
    return out

# map each serviced account to a team: prefer the master 거래처 List mapping; fall back to
# whichever team most of the account's assignees belong to (covers accounts newly added in
# the updated raw data that aren't in the old master account list yet).
account_person_ctr = defaultdict(Counter)
for row in all_rows:
    if row[3] == '취소됨':
        continue
    person = row[7]; account = row[12]
    if person and account and person not in EXCLUDED_PEOPLE and person in member_team:
        account_person_ctr[account][member_team[person]] += 1

def account_team_of(acct):
    if acct in account_team_map:
        return account_team_map[acct]
    ctr = account_person_ctr.get(acct)
    if ctr:
        return ctr.most_common(1)[0][0]
    return None

revisit_by_team = defaultdict(dict)
for acct in account_recs:
    team = account_team_of(acct)
    if team in TEAMS:
        revisit_by_team[team][acct] = account_rates(acct)

print()
print('revisit: accounts with a resolved team:', sum(len(v) for v in revisit_by_team.values()))
for t in TEAMS:
    print(' ', t, 'accounts:', len(revisit_by_team.get(t, {})))

# ---------- validation samples ----------
print()
print('sample: 채대석 이번년도 group counts:', {g: person_group['채대석']['이번 년도'][g] for g in GROUPS_ORDER})
print('sample: 채대석 이번년도 CBC mttr:', person_mttr['채대석']['이번 년도']['CBC'])
print('sample: 허석 이번년도 total case:', sum(person_case['허석']['이번 년도'].values()))
print('sample: 채대석 이번년도 total KPI score:', round(person_score['채대석']['이번 년도'], 2))
print('sample: 채대석 이번년도 score by group:', {g: round(v,2) for g,v in person_score_group['채대석']['이번 년도'].items()})

out = {
    'team_members': team_members,
    'member_team': member_team,
    'fy_meta': {'today': TODAY.isoformat(), 'cur_fy': CUR_FY, 'prev_fy': PREV_FY, 'labels': FY_LABEL},
    'person_group': {p: {per: dict(c) for per, c in pp.items()} for p, pp in person_group.items()},
    'person_mttr': {p: {per: {g: dict(v) for g, v in gg.items()} for per, gg in pp.items()} for p, pp in person_mttr.items()},
    'person_case': {p: {per: dict(c) for per, c in pc.items()} for p, pc in person_case.items()},
    'person_month': {p: {per: {str(m): c for m, c in mc.items()} for per, mc in pm.items()} for p, pm in person_month.items()},
    'person_account_top': {p: c.most_common(6) for p, c in person_account.items()},
    'person_score': {p: dict(d) for p, d in person_score.items()},
    'person_score_group': {p: {per: dict(g) for per, g in d.items()} for p, d in person_score_group.items()},
    'revisit_by_team': revisit_by_team,
}
BUILD_DIR = os.path.join(REPO_ROOT, 'build')
os.makedirs(BUILD_DIR, exist_ok=True)
with open(os.path.join(BUILD_DIR, 'agg_v2.json'), 'w', encoding='utf-8') as f:
    json.dump(out, f, ensure_ascii=False)

print('saved agg_v2.json', os.path.getsize(os.path.join(BUILD_DIR, 'agg_v2.json')), 'bytes')
